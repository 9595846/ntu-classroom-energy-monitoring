from __future__ import annotations

import math
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
OUT_DIR = ROOT.parent / "deliverables"
FIG_DIR = ROOT / "figures"
OUT_DIR.mkdir(exist_ok=True)
FIG_DIR.mkdir(exist_ok=True)
OUTPUT = OUT_DIR / "南通大学教室节能监控方案设计报告.docx"

NAVY = "17365D"
BLUE = "2A6F97"
SKY = "DCEAF4"
PALE = "F4F8FB"
GRID = "D9D9D9"
TEXT = "222222"
MUTED = "5C6773"
ORANGE = "D97706"
GREEN = "2E7D32"

FONT_REG = Path(r"C:\Windows\Fonts\msyh.ttc")
FONT_BOLD = Path(r"C:\Windows\Fonts\msyhbd.ttc")


def pil_font(size: int, bold: bool = False):
    path = FONT_BOLD if bold and FONT_BOLD.exists() else FONT_REG
    if path.exists():
        return ImageFont.truetype(str(path), size=size)
    return ImageFont.load_default()


def text_size(draw: ImageDraw.ImageDraw, text: str, font):
    box = draw.textbbox((0, 0), text, font=font)
    return box[2] - box[0], box[3] - box[1]


def draw_centered(draw, xy, text, font, fill):
    x0, y0, x1, y1 = xy
    w, h = text_size(draw, text, font)
    draw.text(((x0 + x1 - w) / 2, (y0 + y1 - h) / 2 - 2), text, font=font, fill=fill)


def rounded_box(draw, xy, fill, outline, radius=22, width=3):
    draw.rounded_rectangle(xy, radius=radius, fill=fill, outline=outline, width=width)


def arrow(draw, start, end, fill=BLUE, width=6):
    draw.line([start, end], fill=fill, width=width)
    ang = math.atan2(end[1] - start[1], end[0] - start[0])
    head = 18
    for d in (2.55, -2.55):
        p = (end[0] + head * math.cos(ang + d), end[1] + head * math.sin(ang + d))
        draw.line([end, p], fill=fill, width=width)


def make_energy_figure():
    w, h = 1800, 980
    img = Image.new("RGB", (w, h), "white")
    d = ImageDraw.Draw(img)
    title = pil_font(48, True)
    sub = pil_font(28)
    lab = pil_font(30)
    val = pil_font(30, True)
    small = pil_font(25)
    d.text((80, 55), "中位情景下教室年用电与非监管浪费", font=title, fill="#111111")
    d.text((80, 120), "397 间教室  单位 万千瓦时每年", font=sub, fill="#56616D")

    items = [
        ("空调", 254.1, 35.6, "#2A6F97"),
        ("照明", 108.4, 13.0, "#5EA3C7"),
        ("风扇", 8.0, 1.2, "#8ABF8A"),
        ("多媒体", 12.5, 1.0, "#C4A35A"),
    ]
    maxv = 270
    y = 225
    for name, total, waste, color in items:
        d.text((90, y + 20), name, font=lab, fill="#222222")
        x0, x1 = 270, 1570
        d.rounded_rectangle((x0, y, x1, y + 78), radius=16, fill="#ECF1F5")
        tw = (x1 - x0) * total / maxv
        d.rounded_rectangle((x0, y, x0 + tw, y + 78), radius=16, fill=color)
        ww = (x1 - x0) * waste / maxv
        if ww > 4:
            d.rounded_rectangle((x0 + tw - ww, y, x0 + tw, y + 78), radius=16, fill="#D97706")
        d.text((1595, y + 18), f"{total:.1f}", font=val, fill="#222222")
        d.text((x0 + tw - ww - 5, y + 89), f"浪费 {waste:.1f}", font=small, fill="#B85C00")
        y += 165

    d.line((80, 875, 1720, 875), fill="#D7DDE3", width=2)
    d.text((80, 900), "合计年用电约 383.0 万度  非监管浪费约 50.8 万度  占 13.3%", font=pil_font(31, True), fill="#17365D")
    path = FIG_DIR / "energy_baseline.png"
    img.save(path, quality=95)
    return path


def make_architecture_figure():
    w, h = 1900, 1180
    img = Image.new("RGB", (w, h), "white")
    d = ImageDraw.Draw(img)
    title = pil_font(48, True)
    hfont = pil_font(32, True)
    bfont = pil_font(25)
    d.text((75, 45), "教室节能监控系统分层架构", font=title, fill="#111111")

    layers = [
        ("应用层", ["能耗驾驶舱", "告警与工单", "节能考核", "数字孪生"]),
        ("平台层", ["时序数据库", "规则引擎", "模型服务", "统一身份与审计"]),
        ("边缘层", ["楼宇网关", "DDC 本地控制", "断网缓存", "协议转换"]),
        ("设备层", ["远传电表", "智能断路器", "环境与占用传感器", "空调照明执行器"]),
    ]
    fills = ["#E9F2F8", "#EDF5F0", "#FFF4E5", "#F4F1FA"]
    outlines = ["#2A6F97", "#3B7A57", "#D97706", "#74549A"]
    y = 155
    for idx, (label, boxes) in enumerate(layers):
        rounded_box(d, (70, y, 1830, y + 205), fills[idx], outlines[idx], radius=26, width=4)
        rounded_box(d, (95, y + 42, 315, y + 162), outlines[idx], outlines[idx], radius=20, width=2)
        draw_centered(d, (95, y + 42, 315, y + 162), label, hfont, "white")
        x = 360
        for box in boxes:
            rounded_box(d, (x, y + 42, x + 330, y + 162), "white", "#B8C4CE", radius=16, width=2)
            draw_centered(d, (x, y + 42, x + 330, y + 162), box, bfont, "#20262D")
            x += 355
        if idx < len(layers) - 1:
            arrow(d, (950, y + 220), (950, y + 275), fill="#7C8792", width=5)
        y += 255

    d.text((80, 1130), "接口原则  计量与控制解耦  开放协议优先  断网时本地策略仍可运行", font=pil_font(27), fill="#56616D")
    path = FIG_DIR / "system_architecture.png"
    img.save(path, quality=95)
    return path


def make_control_figure():
    w, h = 1900, 1050
    img = Image.new("RGB", (w, h), "white")
    d = ImageDraw.Draw(img)
    d.text((70, 45), "单间教室自动控制流程", font=pil_font(48, True), fill="#111111")
    hfont = pil_font(28, True)
    bfont = pil_font(24)

    boxes = [
        ((90, 190, 400, 320), "课表或预约\n即将开始", "#E9F2F8", "#2A6F97"),
        ((530, 190, 840, 320), "占用确认\n人数大于零", "#EDF5F0", "#3B7A57"),
        ((970, 190, 1280, 320), "按环境条件\n启用设备", "#FFF4E5", "#D97706"),
        ((1410, 190, 1720, 320), "课堂运行\n持续监测", "#F4F1FA", "#74549A"),
        ((1410, 560, 1720, 690), "结束或空置\n进入延时", "#FCECEC", "#B14A4A"),
        ((970, 560, 1280, 690), "关闭照明风扇\n空调转停机", "#FCECEC", "#B14A4A"),
        ((530, 560, 840, 690), "核对回路功率\n生成节能记录", "#EDF5F0", "#3B7A57"),
        ((90, 560, 400, 690), "异常则告警\n派发工单", "#FFF4E5", "#D97706"),
    ]
    for xy, text, fill, outline in boxes:
        rounded_box(d, xy, fill, outline, radius=22, width=4)
        draw_centered(d, xy, text, hfont, "#20262D")
    for a, b in [((400, 255), (530, 255)), ((840, 255), (970, 255)), ((1280, 255), (1410, 255)),
                 ((1565, 320), (1565, 560)), ((1410, 625), (1280, 625)), ((970, 625), (840, 625)), ((530, 625), (400, 625))]:
        arrow(d, a, b, fill="#607080", width=5)

    d.text((90, 790), "关键保护逻辑", font=pil_font(31, True), fill="#17365D")
    notes = [
        "无占用仅执行自动关闭  自动开启须同时满足课表或预约与占用确认",
        "空调通过厂家接口或干接点正常停机  不以频繁硬切电源代替控制",
        "消防应急回路和疏散照明不纳入节能切断  现场手动优先于远程命令",
    ]
    y = 850
    for n in notes:
        d.ellipse((100, y + 8, 116, y + 24), fill="#2A6F97")
        d.text((135, y), n, font=bfont, fill="#333333")
        y += 57
    path = FIG_DIR / "control_flow.png"
    img.save(path, quality=95)
    return path


def make_twin_figure():
    w, h = 1800, 980
    img = Image.new("RGB", (w, h), "white")
    d = ImageDraw.Draw(img)
    d.text((70, 45), "数字孪生的数据闭环", font=pil_font(48, True), fill="#111111")
    centers = [(290, 470), (720, 250), (1150, 250), (1510, 470), (1150, 700), (720, 700)]
    labels = ["空间与设备台账", "实时计量", "状态识别", "策略执行", "效果验证", "模型校准"]
    colors = ["#2A6F97", "#3B7A57", "#D97706", "#74549A", "#B14A4A", "#4C7A8B"]
    for i, (cx, cy) in enumerate(centers):
        d.ellipse((cx - 135, cy - 82, cx + 135, cy + 82), fill="white", outline=colors[i], width=5)
        draw_centered(d, (cx - 135, cy - 82, cx + 135, cy + 82), labels[i], pil_font(26, True), "#20262D")
    for i in range(len(centers)):
        a = centers[i]
        b = centers[(i + 1) % len(centers)]
        vx, vy = b[0] - a[0], b[1] - a[1]
        length = math.hypot(vx, vy)
        s = (a[0] + vx * 150 / length, a[1] + vy * 95 / length)
        e = (b[0] - vx * 150 / length, b[1] - vy * 95 / length)
        arrow(d, s, e, fill="#7C8792", width=5)
    rounded_box(d, (710, 425, 1090, 545), "#17365D", "#17365D", radius=24, width=2)
    draw_centered(d, (710, 425, 1090, 545), "一室一档  一表一回路", pil_font(29, True), "white")
    d.text((125, 875), "孪生模型的价值是把教室  设备  传感器  电表和控制规则绑定到同一对象  支持追溯与优化", font=pil_font(27), fill="#56616D")
    path = FIG_DIR / "digital_twin_loop.png"
    img.save(path, quality=95)
    return path


def make_acceptance_figure():
    w, h = 1900, 640
    img = Image.new("RGB", (w, h), "white")
    d = ImageDraw.Draw(img)
    d.text((70, 40), "试点扩展前的五道验收关卡", font=pil_font(44, True), fill="#111111")
    labels = [
        ("数据可信", "完整率与表计校验"),
        ("控制有效", "命令反馈与功率复核"),
        ("故障可降级", "断网与传感器失联"),
        ("安全合格", "消防隔离与权限审计"),
        ("节能达标", "课时天气归一化"),
    ]
    colors = ["#2A6F97", "#3B7A57", "#D97706", "#74549A", "#B14A4A"]
    x = 70
    for i, (title, sub) in enumerate(labels):
        xy = (x, 190, x + 300, 415)
        rounded_box(d, xy, "white", colors[i], radius=24, width=5)
        draw_centered(d, (x, 215, x + 300, 305), title, pil_font(29, True), "#20262D")
        draw_centered(d, (x + 15, 310, x + 285, 390), sub, pil_font(22), "#56616D")
        if i < len(labels) - 1:
            arrow(d, (x + 315, 302), (x + 360, 302), fill="#7C8792", width=5)
        x += 370
    d.text((75, 520), "任一关卡未通过  先修正试点  不直接扩大到全校", font=pil_font(28, True), fill="#17365D")
    path = FIG_DIR / "acceptance_gates.png"
    img.save(path, quality=95)
    return path


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color=GRID, size="8"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        el = borders.find(qn(tag))
        if el is None:
            el = OxmlElement(tag)
            borders.append(el)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), size)
        el.set(qn("w:color"), color)


def set_cell_margin(cell, top=110, start=120, bottom=110, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn("w:" + m))
        if node is None:
            node = OxmlElement("w:" + m)
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_cant_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    cant_split.set(qn("w:val"), "true")
    tr_pr.append(cant_split)


def set_fixed_width(cell, width_cm):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(int(Cm(width_cm).emu / 635)))
    tc_w.set(qn("w:type"), "dxa")


def set_run_font(run, name="宋体", size=10.5, bold=None, color=TEXT):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Arial")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Arial")
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def set_paragraph_font(paragraph, name="宋体", size=10.5, color=TEXT):
    for run in paragraph.runs:
        set_run_font(run, name, size, color=color)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr_text)
    run._r.append(fld_char2)
    set_run_font(run, "微软雅黑", 9, color=MUTED)


def add_text(doc, text, bold_lead=None, first_line=True, space_after=5, size=10.5):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    p.paragraph_format.space_after = Pt(space_after)
    if first_line:
        p.paragraph_format.first_line_indent = Pt(size * 2)
    if bold_lead and text.startswith(bold_lead):
        r1 = p.add_run(bold_lead)
        set_run_font(r1, "宋体", size, True)
        r2 = p.add_run(text[len(bold_lead):])
        set_run_font(r2, "宋体", size)
    else:
        r = p.add_run(text)
        set_run_font(r, "宋体", size)
    return p


def add_bullet(doc, text, level=0):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.left_indent = Cm(0.7 + level * 0.5)
    p.paragraph_format.first_line_indent = Cm(-0.35)
    p.paragraph_format.line_spacing = 1.25
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(text)
    set_run_font(r, "宋体", 10.3)
    return p


def add_number(doc, text):
    p = doc.add_paragraph(style="List Number")
    p.paragraph_format.left_indent = Cm(0.7)
    p.paragraph_format.first_line_indent = Cm(-0.35)
    p.paragraph_format.line_spacing = 1.25
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(text)
    set_run_font(r, "宋体", 10.3)
    return p


def heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = Pt(10 if level == 1 else 7)
    p.paragraph_format.space_after = Pt(5)
    r = p.add_run(text)
    set_run_font(r, "微软雅黑", 15 if level == 1 else 12, True, "000000")
    return p


def add_caption(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(7)
    r = p.add_run(text)
    set_run_font(r, "宋体", 9.2, color=MUTED)
    return p


def add_figure(doc, path: Path, caption: str, width_inches=6.65):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    shape = p.add_run().add_picture(str(path), width=Inches(width_inches))
    shape._inline.docPr.set("descr", caption)
    shape._inline.docPr.set("title", caption)
    add_caption(doc, caption)


def add_table(doc, headers, rows, widths=None, font_size=9.2, aligns=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.rows[0]._tr.get_or_add_trPr()
    set_repeat_table_header(table.rows[0])
    set_cant_split(table.rows[0])
    for j, h in enumerate(headers):
        c = table.rows[0].cells[j]
        c.text = ""
        set_cell_shading(c, NAVY)
        set_cell_border(c)
        set_cell_margin(c, 120, 120, 120, 120)
        c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        if widths:
            set_fixed_width(c, widths[j])
        p = c.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(str(h))
        set_run_font(r, "微软雅黑", font_size, True, "FFFFFF")
    for i, row in enumerate(rows):
        table_row = table.add_row()
        set_cant_split(table_row)
        cells = table_row.cells
        for j, value in enumerate(row):
            c = cells[j]
            c.text = ""
            if i % 2 == 1:
                set_cell_shading(c, PALE)
            set_cell_border(c)
            set_cell_margin(c)
            c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if widths:
                set_fixed_width(c, widths[j])
            p = c.paragraphs[0]
            p.alignment = (aligns[j] if aligns else WD_ALIGN_PARAGRAPH.LEFT)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.15
            r = p.add_run(str(value))
            set_run_font(r, "宋体", font_size)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def new_page(doc):
    doc.add_page_break()


def configure_document(doc):
    sec = doc.sections[0]
    sec.page_width = Inches(8.5)
    sec.page_height = Inches(11)
    sec.top_margin = Cm(1.85)
    sec.bottom_margin = Cm(1.75)
    sec.left_margin = Cm(2.05)
    sec.right_margin = Cm(2.05)
    sec.header_distance = Cm(0.75)
    sec.footer_distance = Cm(0.75)
    sec.different_first_page_header_footer = True

    normal = doc.styles["Normal"]
    normal.font.name = "宋体"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(TEXT)

    title = doc.styles["Title"]
    title.font.name = "微软雅黑"
    title._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    title.font.color.rgb = RGBColor(0, 0, 0)
    title.font.size = Pt(24)
    title.font.bold = True

    for level, size in ((1, 15), (2, 12), (3, 10.8)):
        s = doc.styles[f"Heading {level}"]
        s.font.name = "微软雅黑"
        s._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
        s.font.color.rgb = RGBColor(0, 0, 0)
        s.font.size = Pt(size)
        s.font.bold = True

    header = sec.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    rr = header.add_run("南通大学教室节能监控方案设计报告")
    set_run_font(rr, "微软雅黑", 8.5, color=MUTED)
    add_page_number(sec.footer.paragraphs[0])


def set_document_properties(doc):
    props = doc.core_properties
    props.title = "南通大学教室节能监控方案设计报告"
    props.subject = "教室节能 自动化监控 数字孪生 远传电表 DDC 智能断路器"
    props.author = ""
    props.keywords = "智能教室,节能,监控,数字孪生,远传电表,DDC,断路器"
    props.comments = "课程作业报告"


def add_cover(doc):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(65)
    r = p.add_run("课程作业一")
    set_run_font(r, "微软雅黑", 14, True, NAVY)

    p = doc.add_paragraph(style="Title")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(18)
    r = p.add_run("南通大学教室节能监控方案设计报告")
    set_run_font(r, "微软雅黑", 24, True, "000000")

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(80)
    r = p.add_run("面向管理 自动化和信息化的一体化设计")
    set_run_font(r, "宋体", 13, color=MUTED)

    meta = [
        ("课程名称", "____________________________"),
        ("学院专业", "____________________________"),
        ("学生姓名", "____________________________"),
        ("学号", "____________________________"),
        ("指导教师", "____________________________"),
    ]
    table = doc.add_table(rows=len(meta), cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for i, (k, v) in enumerate(meta):
        c0, c1 = table.rows[i].cells
        set_fixed_width(c0, 3.0)
        set_fixed_width(c1, 8.0)
        for c in (c0, c1):
            set_cell_margin(c, 120, 100, 120, 100)
            c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        c0.text = ""
        c1.text = ""
        p0 = c0.paragraphs[0]
        p0.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        rr = p0.add_run(k)
        set_run_font(rr, "微软雅黑", 11, True)
        rr = c1.paragraphs[0].add_run(v)
        set_run_font(rr, "宋体", 11, color=MUTED)
        # keep cover table visually open
        for c in (c0, c1):
            tc_pr = c._tc.get_or_add_tcPr()
            borders = OxmlElement("w:tcBorders")
            for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
                el = OxmlElement("w:" + edge)
                el.set(qn("w:val"), "nil")
                borders.append(el)
            tc_pr.append(borders)

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(95)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("二〇二六年十月")
    set_run_font(r, "宋体", 11)


def add_summary(doc):
    heading(doc, "摘要", 1)
    add_text(doc, "本报告针对数百间教室在上课和自习期间缺少持续监管、照明风扇空调可能空转的问题，提出一套以远传电表为计量基础、以 DDC 和设备接口为控制核心、以数字孪生平台为管理载体的节能监控方案。方案保留现场手动控制，自动化系统负责识别空置、执行延时关闭、限制空调设定值、发现异常用电并形成工单。消防应急回路不参与节能切断，空调采用正常停机接口，不采用频繁硬断电方式。")
    add_text(doc, "按南通大学公开资料中的 371 间公共多媒体教室和 26 间公共智慧教室计算，保守纳入 397 间教室。中位情景估算年用电量约 383.0 万千瓦时，其中因空置运行、课表与设备状态不一致、设定值不合理等造成的非监管浪费约 50.8 万千瓦时，占 13.3%。按 0.85 元每千瓦时的测算电价，年浪费电费约 43.2 万元；按 2023 年全国电力平均二氧化碳排放因子 0.5306 千克每千瓦时折算，对应约 269.5 吨二氧化碳。上述结果是工程估算，不替代学校电表实测和财务结算数据。")
    add_text(doc, "建议先选取 30 间不同类型教室开展一学期试点，使用改造前后同类时段和相似天气条件进行校正比较。试点达到数据完整率不低于 99%、误关机事件低于每千教室日 1 次、归一化节电率不低于 12% 后，再按高浪费楼栋优先扩展。全校目标节电量为每年 45 万至 65 万千瓦时。", bold_lead="建议")

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run("关键词  ")
    set_run_font(r, "微软雅黑", 10.5, True)
    r = p.add_run("智能教室  节能监控  数字孪生  远传电表  DDC  智能断路器")
    set_run_font(r, "宋体", 10.5)


def add_toc(doc):
    heading(doc, "目录", 1)
    toc_items = [
        "摘要",
        "一 建设背景与设计边界",
        "二 年度用电量与非监管浪费估算",
        "三 建设目标与总体要求",
        "四 系统总体架构",
        "五 教室自动控制策略",
        "六 信息平台与数字孪生",
        "七 管理机制与运行制度",
        "八 电气安全与网络安全",
        "九 分阶段实施与验收",
        "十 投资与效益分析",
        "十一 结论",
        "附录 A 关键点位表",
        "附录 B 验收测试用例",
        "参考文献",
    ]
    toc_pages = ["2", "4", "5", "7", "8", "10", "12", "13", "14", "15", "17", "18", "19", "20", "21"]
    rows = list(zip(toc_items, toc_pages))
    table = add_table(doc, ["章节", "页码"], rows, widths=[12.0, 2.2], font_size=9.5,
                      aligns=[WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER])
    table._element.set(qn("w:tblCaption"), "目录")


def build():
    fig_energy = make_energy_figure()
    fig_arch = make_architecture_figure()
    fig_control = make_control_figure()
    fig_twin = make_twin_figure()
    fig_acceptance = make_acceptance_figure()

    doc = Document()
    configure_document(doc)
    set_document_properties(doc)
    add_cover(doc)
    new_page(doc)
    add_summary(doc)
    new_page(doc)
    add_toc(doc)
    new_page(doc)

    heading(doc, "一 建设背景与设计边界", 1)
    heading(doc, "1 校园与教室基线", 2)
    add_text(doc, "南通大学现有啬园、启秀和启东 3 个校区。学校公开简介显示，校园建筑面积已超过 130 万平方米，在校本科生 5 万余人。2022 至 2023 学年本科教学质量报告列出公共多媒体教室 371 间、公共智慧教室 26 间，并同时列出语音室、计算机公共实验室和录播教室等其他教学空间[1][2]。2024 年学校又新建或改造智慧教室 8 间、多媒体教室 22 间[3]。由于“改造”可能与原有 397 间重合，本报告不把这 30 间直接累加，而以 397 间作为保守基数，并要求实施前完成逐间台账核验。")
    add_text(doc, "公开资料没有披露逐间教室的设备铭牌功率、课表占用时长和分项电量。本报告因此采用设备功率乘运行时间的自下而上方法估算，只用于论证项目必要性、确定量级和设计试点。正式立项应以配电图、铭牌、课表和连续不少于 4 周的电表数据修正参数。")
    heading(doc, "2 现状问题", 2)
    add_bullet(doc, "管理覆盖存在时间空档。管理员晚间巡查能够处理长明灯和未关空调，但白天课间、临时调课、晚自习结束和节假日临时使用仍依赖使用者自觉。")
    add_bullet(doc, "教室状态分散。课表、门禁或占用、环境参数、设备开关和回路功率往往属于不同系统，无法自动判断“有课但无人”“无人但有功率”等异常。")
    add_bullet(doc, "控制方式不统一。照明、风扇、分体空调和多媒体设备的接口不同，简单远程断电可能损害压缩机，也可能影响教学与安全。")
    add_bullet(doc, "节能责任难量化。若只有楼栋总表，管理部门难以比较同类教室、核实节能措施和定位故障。")
    heading(doc, "3 设计边界", 2)
    add_text(doc, "本方案覆盖普通多媒体教室、智慧教室和开放自习教室中的照明、普通风扇、舒适性空调及教学多媒体非消防负荷。实验室专用设备、数据中心、消防设备、应急照明、电梯和医疗负荷不纳入自动节能切断。系统使用内部管理级计量，不作为供电结算表。")

    new_page(doc)
    heading(doc, "二 年度用电量与非监管浪费估算", 1)
    heading(doc, "1 中位情景参数", 2)
    add_text(doc, "以 397 间教室为基数，假设普通教室平均配置照明 1.30 千瓦、风扇 0.45 千瓦、空调平均输入功率 8.00 千瓦、多媒体设备 0.35 千瓦。运行小时综合考虑教学周、自习开放、夏冬季空调启用和设备非满载。空调开启边界遵循学校现行管理规定，即夏季室外最高温度连续 3 天达到 30 摄氏度以上可制冷，冬季室外最低温度连续 3 天达到 5 摄氏度以下可制热，制冷设定不低于 26 摄氏度、制热不高于 20 摄氏度，并执行人离机停[4]。")
    rows = [
        ("照明", "1.30", "2100", "2730", "12%", "空教室长明和自然采光充足时未分区关闭"),
        ("风扇", "0.45", "450", "202.5", "15%", "课间或离场后未关闭"),
        ("空调", "8.00", "800", "6400", "14%", "空置运行  门窗开启  设定值过低或过高"),
        ("多媒体", "0.35", "900", "315", "8%", "投影与扩声待机或课后未退出"),
    ]
    add_table(doc, ["设备", "平均功率 kW", "年运行 h", "单室年电量 kWh", "非监管比例", "主要原因"], rows,
              widths=[1.8, 2.1, 1.7, 2.3, 1.8, 5.2], font_size=8.7,
              aligns=[WD_ALIGN_PARAGRAPH.CENTER]*5 + [WD_ALIGN_PARAGRAPH.LEFT])
    add_text(doc, "计算关系为 年用电量等于教室数量乘设备平均功率乘年运行小时。非监管浪费等于各类设备年用电量乘相应非监管比例。参数按平均教室折算，空调功率指运行期平均输入功率，不等于铭牌最大值。", first_line=False)
    heading(doc, "2 估算结果", 2)
    rows = [
        ("照明", "108.38", "13.01", "11.1"),
        ("风扇", "8.04", "1.21", "1.0"),
        ("空调", "254.08", "35.57", "30.2"),
        ("多媒体", "12.51", "1.00", "0.9"),
        ("合计", "383.01", "50.78", "43.2"),
    ]
    add_table(doc, ["分项", "年用电量 万kWh", "非监管浪费 万kWh", "浪费电费 万元"], rows,
              widths=[3.0, 3.6, 3.8, 3.3], font_size=9.2,
              aligns=[WD_ALIGN_PARAGRAPH.CENTER]*4)
    add_text(doc, "中位情景下，教室年用电量约 383.0 万千瓦时，非监管浪费约 50.8 万千瓦时。按测算电价 0.85 元每千瓦时，浪费电费约 43.2 万元。0.85 元仅为方案比较用的综合假设值，正式财务测算应替换为学校合同电价、容量电费和分时电价的实际结构。")
    add_figure(doc, fig_energy, "图 1  中位情景下教室年用电构成与非监管浪费  橙色为浪费部分", 6.65)

    heading(doc, "3 不确定性与敏感性", 2)
    rows = [
        ("低情景", "265", "27", "10.2%", "设备功率较低  空调季较短  使用者关机习惯较好"),
        ("中位情景", "383", "51", "13.3%", "采用本报告参数"),
        ("高情景", "530", "86", "16.2%", "大教室比例较高  空调时长较长  调课与自习更频繁"),
    ]
    add_table(doc, ["情景", "年用电 万kWh", "浪费 万kWh", "浪费率", "条件"], rows,
              widths=[2.2, 2.7, 2.5, 2.1, 6.1], font_size=9.2,
              aligns=[WD_ALIGN_PARAGRAPH.CENTER]*4 + [WD_ALIGN_PARAGRAPH.LEFT])
    add_text(doc, "结果对空调平均功率、空调季运行小时和空置运行比例最敏感。中位情景中，空调贡献约 70% 的非监管浪费，因此项目不能只做照明人体感应，还应接入空调的启停和设定值控制。上线后应用远传电表对估算进行校准，形成按校区、楼栋、楼层、教室和设备分项逐级汇总的实测基线。")
    heading(doc, "4 建设必要性", 2)
    add_text(doc, "第一，日常巡查属于事后纠正，无法覆盖数百间教室在全天多个时间窗口中的状态变化。自动化系统可以把空置识别和延时关机下沉到每间教室。第二，远传计量使节能效果可以核证，避免只凭总电费判断。教育部和住房城乡建设部发布的高校校园建筑节能监管导则要求按建筑和空调、动力、照明等分项采集能耗，并设置数据网关、管理平台和运行制度[5][6]。第三，系统同时提高电气异常发现能力，能识别夜间持续功率、回路过载、控制失败和设备效率下降。")
    add_text(doc, "按 2023 年全国电力平均二氧化碳排放因子 0.5306 千克每千瓦时计算，50.8 万千瓦时浪费对应约 269.5 吨二氧化碳[10]。若方案每年节电 45 万至 65 万千瓦时，可减少约 238.8 至 344.9 吨二氧化碳。")

    heading(doc, "三 建设目标与总体要求", 1)
    add_text(doc, "建设目标是形成可计量、可判断、可控制、可追溯的教室用能闭环，并在不影响正常教学和电气安全的前提下降低无效运行。")
    add_number(doc, "建立一室一档。每间教室绑定空间编码、配电回路、电表、设备、传感器、控制器、课表和责任人。")
    add_number(doc, "实现分项计量。至少区分照明插座、空调和多媒体回路；条件不足的旧楼先做到楼层和教室总电量，改造时逐步细分。")
    add_number(doc, "实现本地自治。网络或平台故障时，DDC 仍按课表缓存、占用状态和安全延时运行，现场开关保持可用。")
    add_number(doc, "实现数据驱动管理。平台提供异常告警、工单、同类教室对标、节能量核算和月度公示。")
    add_number(doc, "坚持开放兼容。计量侧优先采用 RS-485 与 Modbus，楼宇侧可使用 BACnet/IP，平台消息接口采用经学校批准的标准协议，避免被单一厂商锁定。")

    new_page(doc)
    heading(doc, "四 系统总体架构", 1)
    add_figure(doc, fig_arch, "图 2  教室节能监控系统分层架构", 6.75)
    heading(doc, "1 设备层", 2)
    add_text(doc, "远传电表采集有功电量、功率、电流、电压和功率因数。智能断路器或带通信的保护装置用于状态监测、故障保护和必要的远程隔离。频繁启停由接触器、调光模块或设备原生控制接口执行，不把断路器当作普通开关。空调优先使用厂家协议网关、集中控制接口或干接点；确需红外控制时，必须增加状态反馈，避免只有命令没有结果。")
    add_text(doc, "占用传感器宜采用被动红外与微波复合方式，大教室可增加毫米波或低分辨率热阵列。照度、温湿度和二氧化碳传感器用于照明分区、空调运行和舒适度验证。默认不使用人脸识别；如利用现有视频估算人数，应在边缘侧只输出人数和置信度，不上传或长期保存人脸图像。")
    heading(doc, "2 边缘层", 2)
    add_text(doc, "每栋教学楼设置一至多个数据网关，完成 Modbus RTU、Modbus TCP、BACnet/IP 等协议转换、时间同步、断网缓存和批量上送。DDC 保存至少 7 天课表与本地策略，平台不可用时仍能执行自动关闭和安全联锁。网关恢复连接后补传数据，不覆盖原始时间戳。GB/T 19582.2-2008 规定了 Modbus 串行链路实现，可作为现场总线设计参考[9]。")
    heading(doc, "3 平台层与应用层", 2)
    add_text(doc, "平台从教务课表、教室预约、校园统一身份和后勤工单系统获取必要数据。时序数据库保存原始测点，规则引擎识别空置用电和异常功率，数字孪生维护空间设备关系。驾驶舱按校区、楼栋和教室逐级下钻，管理人员可查看能耗、设备状态、告警和节能量，但远程控制必须经过角色权限和操作审计。")

    heading(doc, "4 典型设备配置", 2)
    rows = [
        ("教室配电箱", "多功能远传电表  智能保护装置  接触器", "电量功率  保护状态  回路控制", "1 套每室"),
        ("教室顶棚", "复合占用传感器  照度传感器", "占用  光照", "1 至 3 组每室"),
        ("教室墙面", "温湿度与二氧化碳传感器", "环境与舒适度", "1 组每室"),
        ("空调系统", "协议网关或干接点模块", "启停  模式  温度  故障", "按空调系统配置"),
        ("弱电间", "DDC 或边缘控制器  网关", "本地逻辑  缓存  协议转换", "每楼层或每楼栋"),
        ("数据中心", "应用服务器  时序库  备份", "监控  分析  工单  报表", "可复用学校云平台"),
    ]
    add_table(doc, ["位置", "设备", "主要数据或动作", "建议数量"], rows,
              widths=[2.4, 5.2, 4.6, 3.0], font_size=9.0,
              aligns=[WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER])
    heading(doc, "5 计量拓扑", 2)
    add_text(doc, "计量采用校区总表、楼栋总表、楼层分表、教室或设备分表的层级结构。每级电量应满足近似平衡，楼栋下级表之和与楼栋总表的差额用于识别未计量负荷、线路损耗和数据缺失。高校节能监管导则把电表、网关、管理平台和数据中心定义为核心组成，并强调分类分项统计、报表、备份和网络维护[5][6]。本方案在此基础上增加教室级控制和数字孪生对象关系。")
    heading(doc, "6 与智能断路器的边界", 2)
    add_text(doc, "智能断路器的首要功能仍是电路保护。远程分合闸只用于授权隔离、非频繁负荷切换和故障处置。照明和风扇的日常启停使用额定寿命满足频繁操作的接触器或继电器；空调使用设备控制接口完成压缩机保护停机。建筑电气改造应符合 GB 55024-2022 等现行强制性工程建设规范，并由具备相应资质的设计和施工单位完成[8]。")

    heading(doc, "五 教室自动控制策略", 1)
    add_figure(doc, fig_control, "图 3  单间教室自动控制流程", 6.75)
    heading(doc, "1 控制优先级", 2)
    add_text(doc, "控制优先级从高到低为消防和电气保护、现场紧急停止、现场手动操作、经授权的远程操作、自动策略。任何自动命令不得绕过保护装置。系统采用“自动关  谨慎开”的原则，自动关闭以确认空置为依据，自动开启需要课表或预约与占用信号共同成立。")

    new_page(doc)
    heading(doc, "2 典型场景规则", 2)
    rows = [
        ("课前", "课表开始前 15 分钟且检测到占用", "照明按照度分区开启  空调在温度越限时预启  多媒体由教师手动启用"),
        ("上课", "课表有效且有人", "维持教学状态  制冷设定不低于 26 摄氏度  制热不高于 20 摄氏度  监测门窗与二氧化碳"),
        ("短课间", "间隔不超过 20 分钟", "照明降档  风扇按占用  空调进入回差或节能模式  避免频繁启停"),
        ("长空档", "无下一场预约或间隔超过 30 分钟", "照明 5 至 10 分钟后关闭  风扇关闭  空调正常停机  多媒体进入待机或关闭"),
        ("自习", "无课表但有人且在开放时段", "允许现场开启  无人延时关闭  超出开放时段需管理员授权"),
        ("闭馆", "到达楼栋关闭时间", "扫描仍有功率的非消防回路  二次确认后关闭并记录  异常生成工单"),
        ("故障", "传感器冲突  通信中断或功率异常", "保持安全状态  禁止自动开启  允许现场控制  上报告警"),
    ]
    add_table(doc, ["场景", "判定条件", "控制动作"], rows,
              widths=[2.0, 5.0, 8.2], font_size=8.8,
              aligns=[WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT])
    heading(doc, "3 误关机防护", 2)
    add_bullet(doc, "占用判定采用至少两类信息组合，例如课表加占用传感器、门禁事件加微动检测，单一传感器不直接触发关键设备停机。")
    add_bullet(doc, "坐姿学习和考试时人体动作较少，大教室优先采用复合或多点传感器，并在试点期记录误判。")
    add_bullet(doc, "关闭前进行声光提示或面板倒计时，师生可按键续时。人工续时需记录，不要求输入个人身份。")
    add_bullet(doc, "空调设置最小运行和最小停机时间，正常停机后再由保护装置隔离，避免压缩机频繁启动。")
    add_bullet(doc, "设备反馈与电表功率交叉验证。若发出关闭命令后功率仍高于阈值，平台不重复高频下发命令，而是告警并生成工单。")
    heading(doc, "4 照明控制", 2)
    add_text(doc, "照明按讲台、窗侧和内侧分区。有人时根据照度只补足需要的区域，无人时延时关闭。美国能源部面向学校的照明指南建议根据空间选择占用或空置控制，并强调传感器应覆盖入口与工作区、避开空调送风口等误触发位置[11]。其资料汇总的教室照明占用控制节能范围为 10% 至 75%，但跨度很大，本报告只按照明电量的 12% 作为可避免浪费，不直接照搬案例上限[12]。")

    new_page(doc)
    heading(doc, "六 信息平台与数字孪生", 1)
    add_figure(doc, fig_twin, "图 4  数字孪生模型的数据闭环", 6.55)
    heading(doc, "1 对象模型", 2)
    add_text(doc, "数字孪生不是三维展示的同义词。本方案首先建立可计算的对象关系：校区包含楼栋，楼栋包含楼层，楼层包含教室；教室关联设备、回路、电表、传感器、课表、控制策略和工单。对象模型使管理人员能从异常电量追溯到具体回路和设备，也能在设备更换后保留历史记录。是否制作三维模型可按实际需要决定，不作为一期节能的前置条件。")
    heading(doc, "2 数据与功能", 2)
    rows = [
        ("实时监控", "功率  电量  电流电压  占用  温湿度  照度  二氧化碳  设备状态", "1 至 5 分钟采集  关键告警秒级"),
        ("异常识别", "无人有功率  非开放时段用电  同类教室偏高  功率骤增  控制无反馈", "规则先行  稳定后引入模型"),
        ("节能分析", "基线  同比环比  按课时归一化  按天气校正  节能量与碳减排", "保留原始数据与算法版本"),
        ("运维管理", "告警确认  工单派发  到场处理  复核关闭  设备寿命", "与后勤工单系统对接"),
        ("信息发布", "楼栋排名  月报  改造前后对比  典型异常", "公开汇总数据  不公开个人轨迹"),
    ]
    add_table(doc, ["模块", "内容", "实施要点"], rows,
              widths=[2.4, 7.4, 5.4], font_size=8.9,
              aligns=[WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT])
    heading(doc, "3 关键算法", 2)
    add_text(doc, "第一阶段使用可解释规则，包括课表时间窗、占用持续时间、功率阈值和同类教室分位数。积累一个完整学年数据后，再训练负荷预测和故障识别模型。模型输出只用于建议或低风险调整，涉及远程断电的动作仍由确定性规则和权限流程控制。")
    add_text(doc, "基线应按课时数、室外温度和开放时长归一化。节能量等于经校正的基线用电减实施后用电，再扣除新增设备自身用电。所有月报注明数据完整率和估算补点比例，避免用缺失数据制造虚假节能。")

    heading(doc, "七 管理机制与运行制度", 1)
    heading(doc, "1 责任分工", 2)
    rows = [
        ("后勤保障部门", "牵头能源管理  制定策略  处理告警  审核节能量  维护台账"),
        ("信息化部门", "网络与平台运行  接口管理  账号权限  日志审计  备份恢复"),
        ("教务部门", "提供课表与调停课数据  明确开放时段  协调考试等特殊场景"),
        ("学院与物业", "现场巡查  人工复核  工单处理  师生反馈  异常升级"),
        ("电气维保单位", "定期校验  保护试验  设备维修  施工档案  安全培训"),
        ("审计或节能工作组", "抽查数据  审核基线  评价年度目标  监督采购与验收"),
    ]
    add_table(doc, ["责任主体", "主要职责"], rows, widths=[3.6, 11.6], font_size=9.2,
              aligns=[WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT])
    heading(doc, "2 制度安排", 2)
    add_bullet(doc, "建立月度能耗例会。查看数据完整率、异常前十、工单关闭率和同类教室单位课时电耗。")
    add_bullet(doc, "设置策略变更审批。控制阈值、开放时间和批量命令必须留存修改人、原因、生效范围和回退版本。")
    add_bullet(doc, "保留夜间巡查但改变重点。管理员由逐室关灯转为复核平台异常、处理无法自动关闭的设备和现场安全问题。")
    add_bullet(doc, "将节能评价与教学保障并列。若节能率提高但投诉或误关机上升，应先修正策略，不以牺牲教学体验换取数据。")
    add_bullet(doc, "每学年复核教室清单和设备变更，寒暑假执行专门策略，重大活动采用临时授权并在结束后自动失效。")

    new_page(doc)
    heading(doc, "八 电气安全与网络安全", 1)
    heading(doc, "1 电气安全", 2)
    add_text(doc, "设计和施工必须由具备资质的单位完成，校核原配电箱容量、导线截面、短路保护、剩余电流保护、接地和选择性配合。新增控制器不得削弱原有保护功能。消防应急照明、疏散指示和其他消防负荷应通过回路标识、权限和逻辑三重方式排除在节能控制之外。GB 55024-2022 为现行强制性工程建设规范，相关改造应以其及配套现行标准为最低要求[8]。")
    add_text(doc, "控制柜应设置手动自动切换和急停或旁路措施。智能断路器的远程分闸功能默认关闭，仅对通过风险评估的非消防回路开放。批量控制采用分组、限流和间隔启动，避免多个空调同时启动造成冲击电流。")
    heading(doc, "2 网络与数据安全", 2)
    add_bullet(doc, "能源控制网络与办公网逻辑隔离，网关只开放必要的出站连接；现场设备不直接暴露在互联网。")
    add_bullet(doc, "使用学校统一身份认证和基于角色的权限，远程控制、策略修改和数据导出实行最小权限。")
    add_bullet(doc, "关键命令采用双人复核或二次确认，保存操作日志、设备日志和策略版本，日志留存期由学校制度确定。")
    add_bullet(doc, "网关与平台通信加密，设备证书和密钥定期轮换；默认口令必须在安装时更改。")
    add_bullet(doc, "平台按 GB/T 22239-2019 的要求开展定级、差距分析和相应安全建设，具体等级由学校依法确定[7]。")
    add_bullet(doc, "只收集节能所必需的数据。占用信息保存为房间级状态或人数，不形成学生个人轨迹。")
    heading(doc, "3 故障降级", 2)
    rows = [
        ("平台故障", "DDC 按本地缓存策略运行  现场控制不受影响  网关缓存数据"),
        ("网络中断", "禁止新增远程命令  保持本地时钟  恢复后补传"),
        ("传感器故障", "退出相关自动开启  延长关闭确认  提醒人工巡查"),
        ("电表异常", "以保护装置状态为准  停止节能量核算  创建校验工单"),
        ("控制失败", "停止重复动作  保持安全状态  上报告警并现场处置"),
    ]
    add_table(doc, ["故障", "降级策略"], rows, widths=[3.2, 12.0], font_size=9.2,
              aligns=[WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT])

    new_page(doc)
    heading(doc, "九 分阶段实施与验收", 1)
    heading(doc, "1 实施路线", 2)
    rows = [
        ("阶段零", "4 至 6 周", "核对 397 间教室清单  配电图  设备铭牌  网络条件  课表接口  连续测量现状"),
        ("阶段一", "1 个学期", "选择 30 间试点  覆盖新旧楼  大中小教室  空调类型和自习场景  采用对照组"),
        ("阶段二", "4 至 6 个月", "优先改造非监管电量高  投诉少且接口清晰的 100 至 150 间  完成平台和工单对接"),
        ("阶段三", "6 至 12 个月", "扩展至其余教室  校准模型  形成年度审计和持续优化机制"),
    ]
    add_table(doc, ["阶段", "周期", "主要工作"], rows, widths=[2.0, 2.8, 10.4], font_size=9.1,
              aligns=[WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT])
    heading(doc, "2 试点选点", 2)
    add_text(doc, "试点建议在 3 个校区分别选取样本，至少包含 10 间普通多媒体教室、8 间智慧教室、6 间晚自习教室和 6 间大教室。每类设置未改造对照组。试点前记录 4 周基线，覆盖空调季更佳；若学期时间不足，可先完成计量，自动控制在下一空调季验收。")

    heading(doc, "3 验收指标", 2)
    rows = [
        ("数据完整率", "不低于 99%", "按 5 分钟测点统计  计划停机单列"),
        ("电量准确性", "与校验表偏差满足选用表计等级和项目要求", "逐表校验并进行上下级平衡"),
        ("命令成功率", "不低于 99%", "平台命令有执行反馈和功率复核"),
        ("误关机", "低于每千教室日 1 次", "以师生投诉和现场复核为准"),
        ("告警闭环率", "不低于 95%", "24 小时内确认  按等级限时关闭"),
        ("节电率", "试点归一化后不低于 12%", "同类对照  课时和天气校正"),
        ("本地自治", "平台与网络中断时策略可运行", "断网 2 小时实测"),
        ("安全", "消防回路无误控  保护功能无削弱", "专项测试和文件审查"),
    ]
    add_table(doc, ["指标", "建议门槛", "核验方法"], rows, widths=[3.0, 4.0, 8.2], font_size=8.9,
              aligns=[WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT])
    heading(doc, "4 测量与验证", 2)
    add_text(doc, "验收不只比较总电量。应按课时、开放时长、室外温度和教室类型校正，并报告统计区间。对改造前后采用相同房间比较时，应排除大型考试、停课和施工；采用对照组时，应说明对照组与试点组的容量和使用差异。节能量计算文件、原始数据、参数版本和异常剔除记录一并归档。")
    heading(doc, "5 验收资料", 2)
    rows = [
        ("竣工与资产", "竣工图  设备清单  编码标签  合格证  质保文件"),
        ("数据与接口", "点位表  数据字典  接口文档  账号权限矩阵"),
        ("测试与校准", "电表校验  控制测试  断网测试  安全专项记录"),
        ("软件与安全", "配置备份  策略版本  漏洞修复记录  恢复预案"),
        ("运行与培训", "操作手册  培训签到  工单流程  维保联系人"),
    ]
    add_table(doc, ["资料类别", "交付内容"], rows, widths=[3.3, 11.9], font_size=9.0,
              aligns=[WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT])
    add_figure(doc, fig_acceptance, "图 5  试点扩展前的验收关卡", 6.55)

    new_page(doc)
    heading(doc, "十 投资与效益分析", 1)
    heading(doc, "1 投资估算", 2)
    rows = [
        ("教室分项计量", "397 间", "约 48", "远传电表  电流互感器  辅材"),
        ("传感器与 DDC", "397 间", "约 95", "占用  照度  温湿度  控制器"),
        ("空调接口与现场旁路", "397 间", "约 36", "协议模块  干接点  面板"),
        ("网关与网络", "按楼栋", "约 22", "边缘网关  交换与安装"),
        ("平台与数字孪生", "1 套", "约 62", "软件  数据库  接口  可视化"),
        ("安装调试", "全校", "约 52", "配电改造  测试  标签  竣工资料"),
        ("培训备件与预备费", "全校", "约 51", "培训  备件  约 12% 预备费"),
        ("合计", "", "约 366", "合理区间约 320 至 450 万元"),
    ]
    add_table(doc, ["项目", "规模", "估算 万元", "说明"], rows,
              widths=[4.0, 2.4, 2.8, 6.0], font_size=8.9,
              aligns=[WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT])
    add_text(doc, "投资估算用于方案比选，不替代施工图预算和招标控制价。若智慧教室已有占用传感器、控制器和网络，可复用并减少投资；旧楼配电箱空间不足、空调品牌接口不统一或需重新布线时，投资会上升。建议把平台、网关和教室节点分项报价，明确协议开放、数据归属和后续扩容单价。")

    heading(doc, "2 直接节能收益", 2)
    add_text(doc, "全校运行稳定后，建议把年度节电目标设为 45 万至 65 万千瓦时，相当于中位基线的 11.7% 至 17.0%。按 0.85 元每千瓦时计算，年直接电费收益约 38.3 万至 55.3 万元。按 320 万至 450 万元投资计算，单纯以电费衡量的静态回收期约 5.8 至 11.8 年。中位投资 366 万元和中位年收益 46.8 万元对应约 7.8 年。")
    add_text(doc, "直接收益没有计入减少人工巡查、提前发现空调和回路故障、降低设备无效运行小时、辅助空间利用分析等价值，也没有计入平台自身电耗和年度维护费。正式可研应分别列出这些项目，避免重复计算。")
    heading(doc, "3 方案分级", 2)
    rows = [
        ("基础计量", "楼栋楼层与重点教室远传计量  异常告警", "投资较低  快速找到浪费点", "不能自动消除全部浪费"),
        ("推荐方案", "397 间计量  占用感知  DDC 控制  平台与工单", "形成完整闭环  节能可核证", "需要跨部门协同和调试"),
        ("深度改造", "推荐方案加 LED  调光  空调系统改造  新风优化", "节能潜力更高  改善舒适度", "投资与施工影响较大"),
    ]
    add_table(doc, ["级别", "范围", "优点", "局限"], rows,
              widths=[2.3, 5.3, 4.0, 3.6], font_size=8.9,
              aligns=[WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT])
    heading(doc, "4 决策建议", 2)
    add_text(doc, "先批准计量和 30 间试点，再以实测数据决定全量投资。若试点显示主要浪费集中在少数楼栋，可采用 80 20 策略优先覆盖；若空调接口改造成本过高，可先完成计量、设定值管理和离场提醒，待设备更新时接入原生控制。采购合同应要求开放数据字典、接口文档、配置备份和迁移能力。")

    heading(doc, "十一 结论", 1)
    add_text(doc, "南通大学教室数量多、开放时段长，依靠晚间人工巡查无法覆盖全天非监管时段。基于公开教室数量和工程参数的中位估算表明，397 间教室年用电约 383.0 万千瓦时，非监管浪费约 50.8 万千瓦时。空调是最大来源，应与照明、风扇和多媒体设备一并纳入控制。")
    add_text(doc, "推荐方案以远传电表提供事实基线，以 DDC 和设备接口执行本地策略，以数字孪生关联空间、设备和数据，以管理制度完成告警和工单闭环。项目应从 30 间多类型教室试点开始，并以数据完整率、误关机率、归一化节电率和安全测试作为扩展条件。这样可以在保障教学的前提下，把节能从依赖个人习惯的巡查工作转变为持续、可核证的运行管理。")

    new_page(doc)
    heading(doc, "附录 A 关键点位表", 1)
    rows = [
        ("电表", "累计有功电量", "kWh", "5 分钟", "计量  基线  节能核算"),
        ("电表", "有功功率", "kW", "1 分钟", "空置用电和控制反馈"),
        ("电表", "电压电流功率因数", "V A", "5 分钟", "电气异常与设备状态"),
        ("断路器", "分合闸与故障状态", "状态", "事件", "保护和隔离"),
        ("占用", "有人无人  人数", "状态 人", "10 秒", "控制条件  使用率"),
        ("环境", "温度湿度", "摄氏度 百分比", "1 分钟", "空调策略和舒适度"),
        ("环境", "照度", "lx", "1 分钟", "照明分区"),
        ("环境", "二氧化碳", "ppm", "1 分钟", "通风提示"),
        ("空调", "模式  设定温度  故障", "状态", "事件", "运行与故障"),
        ("照明风扇", "命令与反馈", "状态", "事件", "控制闭环"),
        ("课表", "课程  起止时间  状态", "记录", "变更时", "使用计划"),
        ("工单", "告警  责任人  处理结果", "记录", "事件", "管理闭环"),
    ]
    add_table(doc, ["对象", "点位", "单位", "建议周期", "用途"], rows,
              widths=[2.0, 4.1, 2.6, 2.3, 4.2], font_size=8.7,
              aligns=[WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT])
    add_text(doc, "点位命名建议采用 校区 楼栋 楼层 房间 系统 设备 点位 的层级编码。原始测点、计算点和人工录入数据使用不同前缀，避免混淆。所有时间戳统一到学校标准时间。", first_line=False)

    new_page(doc)
    heading(doc, "附录 B 验收测试用例", 1)
    rows = [
        ("B01", "有课有人", "模拟课表有效并进入教室", "允许按环境条件启用  状态与功率一致"),
        ("B02", "有课无人", "课表有效但持续无人", "不自动开启  已开设备延时关闭"),
        ("B03", "自习续时", "无课表  人员按键续时", "在开放时段延长  记录人工覆盖"),
        ("B04", "传感器失联", "断开占用传感器", "进入降级  不误关  产生告警"),
        ("B05", "平台断网", "阻断楼宇网关上行 2 小时", "本地策略运行  数据缓存并补传"),
        ("B06", "空调停机", "发送关闭命令", "设备正常停机  满足最小停机时间  功率下降"),
        ("B07", "控制失败", "执行器拒动", "停止重复动作  告警和工单生成"),
        ("B08", "消防回路", "执行闭馆批量命令", "消防与应急回路不受影响"),
        ("B09", "越权操作", "普通用户尝试远程分闸", "拒绝操作并记录审计日志"),
        ("B10", "电量平衡", "汇总下级表与上级表", "差额在项目规定范围内  缺失可追溯"),
    ]
    add_table(doc, ["编号", "场景", "步骤", "预期结果"], rows,
              widths=[1.5, 2.7, 5.3, 5.7], font_size=8.7,
              aligns=[WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT])

    new_page(doc)
    heading(doc, "参考文献", 1)
    refs = [
        "[1] 南通大学. 学校简介. https://www.ntu.edu.cn/xxjj/list.htm  访问日期 2026 10 22.",
        "[2] 南通大学. 南通大学 2022 至 2023 学年本科教学质量报告. https://xxgk.ntu.edu.cn/2023/1205/c5182a223717/page.htm  2023 12 05.",
        "[3] 南通大学. 南通大学 2024 年十件大事. https://news.ntu.edu.cn/_t61/2025/0105/c9a258816/page.htm  2025 01 05.",
        "[4] 南通大学后勤保障部. 南通大学空调器使用管理规定 2020 年修订. https://hqbzb.ntu.edu.cn/2020/1203/c5583a158201/page.htm  2020 05 21.",
        "[5] 住房和城乡建设部  教育部. 高等学校校园建筑节能监管系统建设技术导则及有关管理办法. 建科 2009 163 号  2009.",
        "[6] 高等学校校园建筑节能监管系统运行管理技术导则. https://www.hncu.edu.cn/xjjyhqglc/info/1093/2933.htm  访问日期 2026 10 22.",
        "[7] 国家市场监督管理总局  国家标准化管理委员会. GB/T 22239-2019 信息安全技术 网络安全等级保护基本要求. https://openstd.samr.gov.cn/bzgk/std/newGbInfo?hcno=BAFB47E8874764186BDB7865E8344DAF.",
        "[8] 住房和城乡建设部. GB 55024-2022 建筑电气与智能化通用规范. https://www.mohurd.gov.cn/gongkai/zc/wjk/art/2024/art_17339_765632.html.",
        "[9] 国家标准化管理委员会. GB/T 19582.2-2008 基于 Modbus 协议的工业自动化网络规范 第 2 部分. https://openstd.samr.gov.cn/bzgk/std/newGbInfo?hcno=58A740D39A7E96EE7A3001F5FB73BE65.",
        "[10] 生态环境部  国家统计局. 2023 年电力二氧化碳排放因子. https://www.mee.gov.cn/ywdt/zbft/202601/t20260105_1139911.shtml  2026 01 05.",
        "[11] U.S. Department of Energy. Lighting Specification Guidance for Schools. 2024. https://www.energy.gov/sites/default/files/2024-12/lighting-spec-guidance-school_nov24.pdf.",
        "[12] U.S. Department of Energy Better Buildings. Wireless Occupancy Sensors for Lighting Controls. https://www.govinfo.gov/content/pkg/GOVPUB-E-PURL-gpo177300/pdf/GOVPUB-E-PURL-gpo177300.pdf.",
        "[13] 住房和城乡建设部. GB 55015-2021 建筑节能与可再生能源利用通用规范. https://www.beijing.gov.cn/zhengce/zhengcefagui/qtwj/202204/t20220412_2674410.html.",
    ]
    for ref in refs:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.65)
        p.paragraph_format.first_line_indent = Cm(-0.65)
        p.paragraph_format.space_after = Pt(5)
        p.paragraph_format.line_spacing = 1.15
        r = p.add_run(ref)
        set_run_font(r, "宋体", 9.2)

    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
